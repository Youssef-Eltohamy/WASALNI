# -*- coding: utf-8 -*-
"""Generate an Arabic (RTL) Word doc explaining the 3 DB-schema decisions for WASALNI."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Tahoma"  # ships on all Windows, renders Arabic cleanly
PRIMARY = RGBColor(0x0D, 0x5C, 0x75)      # WASALNI petrol
PRIMARY_DARK = RGBColor(0x08, 0x2F, 0x3D)
ACCENT = RGBColor(0xFF, 0x7A, 0x45)       # coral
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
GREY = RGBColor(0x55, 0x55, 0x55)


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def _rtl_run(run):
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)


def add_par(doc, text="", size=11, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.RIGHT,
            space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    _set_rtl(p)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        _rtl_run(r)
    return p


def add_heading(doc, text, level=1):
    sizes = {0: 22, 1: 16, 2: 13}
    colors = {0: PRIMARY_DARK, 1: PRIMARY, 2: PRIMARY_DARK}
    p = add_par(doc, text, size=sizes.get(level, 12), bold=True,
                color=colors.get(level, PRIMARY), space_before=12, space_after=6)
    return p


def add_bullet(doc, text, color=None, sym="•"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.right_indent = Inches(0.25)
    r = p.add_run(f"{sym}  {text}")
    r.font.name = FONT
    r.font.size = Pt(11)
    if color is not None:
        r.font.color.rgb = color
    _rtl_run(r)
    return p


def add_rule(doc):
    p = add_par(doc, "", space_after=2, space_before=2)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D0DCE2")
    pbdr.append(bottom)
    pPr.append(pbdr)


doc = Document()
# Default style font + RTL section
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
sectPr = doc.sections[0]._sectPr
bidi = OxmlElement("w:bidi")
sectPr.append(bidi)

# ---------- Title ----------
add_heading(doc, "قرارات تصميم قاعدة البيانات — WASALNI", level=0)
add_par(doc, "المرحلة: الـ Data Model (Supabase Schema)  •  الفرع: 002-rebuild-from-zero  •  التاريخ: 2026-05-29",
        size=10, color=GREY, space_after=10)
add_par(doc,
        "الملف ده بيشرح 3 قرارات محتاجة حسم قبل ما نكتب الـ migration الخاص بقاعدة البيانات. "
        "لكل سؤال: الخلفية، الخيارات المتاحة، مميزات وعيوب كل خيار، والتوصية مع السبب.",
        size=11, space_after=8)
add_rule(doc)

# ========== Q1 ==========
add_heading(doc, "السؤال 1: بنية الـ Listings — جدول موحّد ولا جداول منفصلة؟", level=1)
add_par(doc,
        "عندنا 3 أنواع من الكيانات المعروضة في التطبيق: مقدّم خدمة (سباك/كهربائي...)، محل، "
        "وسائق نقل. الـ 3 بيشتركوا في ~80% من البيانات (اسم، نبذة، تليفون، موقع، قرية، فئة، حالة، "
        "توثيق، تمييز). السؤال: نمثّلهم في جدول واحد ولا 3 جداول منفصلة؟",
        space_after=8)

add_heading(doc, "الخيار أ (المقترح): جدول موحّد listings بعمود kind", level=2)
add_par(doc, "جدول واحد اسمه listings فيه عمود kind قيمته (service / shop / transport)، "
             "والحقول الخاصة بكل نوع تتخزّن في عمود attributes من نوع jsonb "
             "(مثلاً نوع المركبة للنقل، أو خدمة التوصيل للمحل).", space_after=6)
add_par(doc, "المميزات:", bold=True, color=GREEN, space_after=3)
for t in [
    "الـ Feed والبحث = استعلام واحد على جدول واحد (أسرع وأبسط).",
    "منطق المفضلة والتوثيق والتمييز (featured) يتكتب مرة واحدة بدل 3 مرات.",
    "إضافة نوع جديد مستقبلاً = صف بيانات، مش جدول جديد + كود جديد.",
    "كود أقل بكتير — مناسب لتطوير فردي (solo dev + Claude).",
    "العلاقات (صور، منتجات، طلبات، توثيق) كلها بتربط بجدول واحد.",
]:
    add_bullet(doc, t, color=GREEN)
add_par(doc, "العيوب:", bold=True, color=ACCENT, space_after=3, space_before=4)
for t in [
    "الحقول الخاصة بنوع معيّن بتروح في attributes (jsonb) — الفلترة عليها أعقد شوية، بس محلولة بـ GIN index.",
    "الصف ممكن يكون فيه أعمدة فاضية حسب النوع (مثلاً cover_url للمحلات بس) — تأثيره بسيط جداً.",
    "التحقق إن كل نوع عنده الحقول الصح بيتعمل على مستوى التطبيق، مش قيود DB صارمة لكل نوع.",
]:
    add_bullet(doc, t, color=ACCENT)

add_heading(doc, "الخيار ب: جداول منفصلة (providers + shops + transport_drivers)", level=2)
add_par(doc, "كل نوع له جدوله الخاص بأعمدته المحددة.", space_after=6)
add_par(doc, "المميزات:", bold=True, color=GREEN, space_after=3)
for t in [
    "أعمدة كل جدول واضحة ومحددة لنوعها — أسهل في القراءة الأولية للـ schema.",
    "قيود قاعدة البيانات أقوى (NOT NULL لكل حقل حسب نوعه).",
]:
    add_bullet(doc, t, color=GREEN)
add_par(doc, "العيوب:", bold=True, color=ACCENT, space_after=3, space_before=4)
for t in [
    "تكرار نفس المنطق 3 مرات (المفضلة، التوثيق، التمييز، البحث) — كود أكتر وأخطاء أكتر.",
    "الـ Feed الموحّد محتاج UNION على 3 جداول → أبطأ وأعقد، والترتيب/الترقيم (pagination) يبقى وجع.",
    "إضافة نوع رابع مستقبلاً = جدول جديد + إعادة كتابة كل الاستعلامات.",
    "صيانة أتقل بكتير على مطوّر فردي.",
]:
    add_bullet(doc, t, color=ACCENT)

add_par(doc, "✅ التوصية: الخيار أ (موحّد).", bold=True, color=PRIMARY, space_before=8, space_after=3)
add_par(doc,
        "السبب: المشروع تطوير فردي (إنت + Claude)، والأنواع الثلاثة بتشترك في معظم السلوك "
        "(feed، بحث، مفضلة، توثيق، تمييز). الجدول الموحّد بيوفّر تكرار ضخم في الكود ويخلّي إضافة "
        "أنواع جديدة سهلة. عيبه الوحيد (الحقول الخاصة في jsonb) مقبول تماماً في حجم الـ MVP ومحلول بالفهارس.",
        space_after=8)
add_rule(doc)

# ========== Q2 ==========
add_heading(doc, "السؤال 2: تخزين صور التوثيق الحسّاسة", level=1)
add_par(doc,
        "عند التوثيق، مقدّم الخدمة/المحل بيرفع صورة بطاقة الرقم القومي + صورة للنشاط/سيلفي. "
        "سياسة الخصوصية بتقول: البيانات دي حسّاسة، لازم تتخزّن مشفّرة وتتحذف بعد المراجعة "
        "(approve/reject). السؤال: نعملها صح من الأول ولا نبسّطها؟",
        space_after=8)

add_heading(doc, "الخيار أ (المقترح): Bucket خاص + حذف تلقائي + تشفير", level=2)
add_par(doc, "Supabase Storage bucket خاص (مش public)، الوصول للصور بـ signed URLs للأدمن فقط، "
             "الصور تتحذف تلقائياً بعد ما الأدمن يعمل approve أو reject، والرقم القومي نفسه "
             "يتخزّن مشفّر بـ pgcrypto.", space_after=6)
add_par(doc, "المميزات:", bold=True, color=GREEN, space_after=3)
for t in [
    "متوافق مع سياسة الخصوصية اللي اتحطّت من الأول.",
    "أمان حقيقي للبيانات الشخصية — مفيش صور بطاقات متخزّنة للأبد.",
    "لو حصل تسريب، مفيش بيانات حسّاسة قديمة معرّضة.",
    "بيبني ثقة المستخدمين (مبدأ \"الثقة\" من مبادئ المشروع الثلاثة).",
]:
    add_bullet(doc, t, color=GREEN)
add_par(doc, "العيوب:", bold=True, color=ACCENT, space_after=3, space_before=4)
for t in [
    "شغل إضافي دلوقتي: إعداد الـ bucket الخاص + سياسات الوصول + آلية الحذف التلقائي.",
    "لو احتجنا نراجع توثيق قديم، الصور هتكون اتحذفت (بس ده المطلوب قانونياً).",
]:
    add_bullet(doc, t, color=ACCENT)

add_heading(doc, "الخيار ب: تبسيطها دلوقتي", level=2)
add_par(doc, "نعمل bucket خاص بس، ونأجّل التشفير والحذف التلقائي لمرحلة لاحقة.", space_after=6)
add_par(doc, "المميزات:", bold=True, color=GREEN, space_after=3)
for t in [
    "أسرع في التنفيذ دلوقتي — نخلّص الـ MVP بدري.",
    "كود أقل في البداية.",
]:
    add_bullet(doc, t, color=GREEN)
add_par(doc, "العيوب:", bold=True, color=ACCENT, space_after=3, space_before=4)
for t in [
    "مخالف لسياسة الخصوصية المعلنة — صور بطاقات متخزّنة بدون حذف.",
    "خطر قانوني/سمعة لو حصل تسريب لبيانات شخصية مصرية.",
    "ديْن تقني: تطبيق التشفير/الحذف لاحقاً على بيانات موجودة أصعب من عمله من الأول.",
]:
    add_bullet(doc, t, color=ACCENT)

add_par(doc, "✅ التوصية: الخيار أ (Private + حذف + تشفير).", bold=True, color=PRIMARY, space_before=8, space_after=3)
add_par(doc,
        "السبب: دي بيانات شخصية حسّاسة (بطاقات رقم قومي)، والسياسة اتقررت من الأول. الشغل الإضافي "
        "بسيط ومحدود، والمخاطرة القانونية والسمعية في الخيار ب أكبر بكتير من الوقت اللي هنوفّره. "
        "كمان \"الثقة\" مبدأ أساسي في WASALNI.",
        space_after=8)
add_rule(doc)

# ========== Q3 ==========
add_heading(doc, "السؤال 3: الخطوة الجاية", level=1)
add_par(doc, "بعد ما تراجع القرارين فوق، نمشي إزاي؟", space_after=8)

add_heading(doc, "الخيار أ: اكتب الـ migration", level=2)
add_par(doc, "أبدأ أكتب ملف الـ SQL migration الكامل (enums + الجداول الـ12 + الفهارس + سياسات "
             "RLS + triggers التحديث التلقائي) وأطبّقه على Supabase.", space_after=4)
add_par(doc, "مناسب لو القرارين فوق متوافقين مع رؤيتك.", color=GREY, size=10, space_after=6)

add_heading(doc, "الخيار ب: عايز تعدّل الأول", level=2)
add_par(doc, "فيه حاجة في الجداول/الأعمدة عايز تغيّرها أو تضيفها قبل ما أكتب الكود "
             "(مثلاً حقول إضافية، جدول ناقص، تسمية مختلفة).", space_after=4)
add_par(doc, "مناسب لو حابب تراجع التفاصيل أو عندك متطلبات مش مغطّاة.", color=GREY, size=10, space_after=6)

add_par(doc, "✅ التوصية: مرتبطة بالقرارين فوق.", bold=True, color=PRIMARY, space_before=8, space_after=3)
add_par(doc,
        "لو موافق على الموحّد + التوثيق الآمن، نمشي بكتابة الـ migration على طول. لو عندك تعديلات، "
        "قولها الأول ونظبّط الـ schema قبل أي كود.",
        space_after=8)
add_rule(doc)

# ========== Appendix ==========
add_heading(doc, "ملحق: ملخص الجداول الـ12", level=1)
tables = [
    ("villages", "القرى (كفر المقدام + تفهنا الأشراف + التوسع)"),
    ("profiles", "كل المستخدمين، يمتد من auth.users (Phone OTP)"),
    ("categories", "الفئات الشجرية (خدمة/محل/نقل)"),
    ("listings", "⭐ الكيان الموحّد: خدمة/محل/نقل (kind + attributes)"),
    ("listing_photos", "معرض الصور / portfolio"),
    ("products", "منتجات المحلات"),
    ("favorites", "المفضلة (listing أو product)"),
    ("order_intents", "⭐ tracking للطلبات (analytics فقط، مفيش payment)"),
    ("verifications", "التوثيق (بطاقة + صورة، مشفّرة، تتحذف بعد المراجعة)"),
    ("admin_users", "السوبر أدمن (منفصل عن profiles للأمان)"),
    ("audit_logs", "سجل عمليات الأدمن"),
    ("app_settings", "إعدادات عامة (تاريخ الإطلاق، flags)"),
]
for name, desc in tables:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.right_indent = Inches(0.2)
    r1 = p.add_run(f"{desc}")
    r1.font.name = FONT
    r1.font.size = Pt(11)
    _rtl_run(r1)
    r2 = p.add_run(f"  —  {name}")
    r2.font.name = "Consolas"
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.color.rgb = PRIMARY

out = r"D:\programing\wasalni\return_to_zero\قرارات_قاعدة_البيانات.docx"
doc.save(out)
print("SAVED:", out)
