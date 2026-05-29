# -*- coding: utf-8 -*-
"""Generate an Arabic (RTL) Word doc explaining the 3 frontend build approaches for WASALNI."""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Tahoma"
PRIMARY = RGBColor(0x0D, 0x5C, 0x75)
PRIMARY_DARK = RGBColor(0x08, 0x2F, 0x3D)
ACCENT = RGBColor(0xFF, 0x7A, 0x45)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
GREY = RGBColor(0x55, 0x55, 0x55)


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))


def _rtl_run(run):
    rPr = run._r.get_or_add_rPr()
    rPr.append(OxmlElement("w:rtl"))


def add_par(doc, text="", size=11, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.RIGHT,
            space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    _set_rtl(p)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        _rtl_run(r)
    return p


def add_heading(doc, text, level=1):
    sizes = {0: 22, 1: 16, 2: 13}
    colors = {0: PRIMARY_DARK, 1: PRIMARY, 2: PRIMARY_DARK}
    return add_par(doc, text, size=sizes.get(level, 12), bold=True,
                   color=colors.get(level, PRIMARY), space_before=12, space_after=6)


def add_bullet(doc, text, color=None, sym="•", indent=0.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.right_indent = Inches(indent)
    r = p.add_run(f"{sym}  {text}")
    r.font.name = FONT
    r.font.size = Pt(11)
    if color is not None:
        r.font.color.rgb = color
    _rtl_run(r)
    return p


def add_rule(doc):
    p = add_par(doc, "", space_after=2, space_before=2)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D0DCE2")
    pbdr.append(bottom)
    pPr.append(pbdr)


def pros(doc, items):
    add_par(doc, "المميزات:", bold=True, color=GREEN, space_after=3, space_before=4)
    for t in items:
        add_bullet(doc, t, color=GREEN, sym="✔")


def cons(doc, items):
    add_par(doc, "العيوب:", bold=True, color=ACCENT, space_after=3, space_before=4)
    for t in items:
        add_bullet(doc, t, color=ACCENT, sym="✕")


def see(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(4)
    r1 = p.add_run("هتشوف إيه:  ")
    r1.font.name = FONT; r1.font.size = Pt(11); r1.font.bold = True; r1.font.color.rgb = PRIMARY
    _rtl_run(r1)
    r2 = p.add_run(text)
    r2.font.name = FONT; r2.font.size = Pt(11)
    _rtl_run(r2)


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
doc.sections[0]._sectPr.append(OxmlElement("w:bidi"))

# ---------- Title ----------
add_heading(doc, "مناهج بناء واجهة WASALNI (Frontend-First)", level=0)
add_par(doc, "المرحلة: تطبيق Flutter — موبايل فقط  •  الفرع: 002-rebuild-from-zero  •  التاريخ: 2026-05-29",
        size=10, color=GREY, space_after=10)
add_par(doc,
        "قرّرنا نبني الواجهة كاملة الأول بـ mock data (بيانات وهمية) قبل ما نشتغل في الباك إند، "
        "عشان لما الـ UI يكون كامل قدامنا نقدر نستخرج المواصفات الدقيقة ونبني الـ schema على أساسها. "
        "الملف ده بيشرح: جرد الشاشات، وتلات مناهج لترتيب البناء بمميزاتها وعيوبها، والتوصية — عشان تختار.",
        space_after=8)
add_rule(doc)

# ---------- Screen inventory ----------
add_heading(doc, "جرد الشاشات — \"الفرونت الكامل\" = ~18 شاشة", level=1)
groups = [
    ("الاكتشاف والتصفّح", ["الرئيسية (Feed) — جريد عمودين + تاب خريطة", "التصنيفات", "البحث", "فلتر القرية (قريتي / المجاورة)"]),
    ("البروفايلات", ["صفحة مقدم خدمة — بورتفوليو", "صفحة محل — منتجات + توصيل", "صفحة منتج", "صفحة سائق نقل"]),
    ("السلة والطلب", ["السلة", "تأكيد الطلب ← تسليم لواتساب"]),
    ("الحساب والدخول", ["onboarding / اختيار القرية", "موبايل + OTP", "استكمال البروفايل", "حسابي", "المفضلة"]),
    ("تسجيل النشاط", ["سجّل نشاطك (إنشاء listing)", "رفع التوثيق"]),
    ("متفرقات", ["الإعدادات / الشروط / الخصوصية"]),
]
for title, items in groups:
    add_par(doc, title, bold=True, color=PRIMARY_DARK, space_after=2, space_before=4)
    for it in items:
        add_bullet(doc, it, indent=0.3)
add_par(doc, "بالإضافة إلى: الـ bottom navigation بـ 4 تابات (الرئيسية / التصنيفات / المفضلة / حسابي).",
        size=10, color=GREY, space_before=6, space_after=4)
add_rule(doc)

# ---------- Approach A ----------
add_heading(doc, "المنهج أ) Depth-first — شاشة كاملة في كل مرة", level=1)
add_par(doc, "الترتيب: الأساس ← الرئيسية بالكامل (جريد + كروت + فلتر + تاب خريطة + Bloc + mock repo + "
             "حالات تحميل/فاضي/خطأ + تلميع كامل) ← بعدها شاشة كاملة تانية ← وهكذا. كل شاشة 100% قبل اللي بعدها.")
see(doc, "بعد أول milestone الرئيسية مظبوطة وحلوة — بس لو دوست على كارت مفيش حاجة تتفتح (الشاشة لسه ماتعملتش). "
        "التطبيق \"عميق وضيّق\": كام شاشة مثالية والباقي غايب. تشوف التطبيق كامل في الآخر خالص.")
pros(doc, ["كل شاشة بتتسلّم بجودة إنتاج كاملة.", "صفر إعادة شغل.", "المعمارية تتأكد من أول شاشة."])
cons(doc, ["ماتعيشش الـ flow الكامل غير قرب النهاية.", "صعب تقيّم الـ navigation والتدفق ككل.",
           "هدفك (تشوف الـ UI كامل) يتحقق آخر حاجة."])
add_rule(doc)

# ---------- Approach B ----------
add_heading(doc, "المنهج ب) Breadth-first — هيكل قابل للتنقّل بالكامل", level=1)
add_par(doc, "الترتيب: الأساس ← كل الـ18 شاشة كـ stubs خشنة لكن كلها شغّالة ومترابطة (الـ bottom nav يشتغل، "
             "دوس كارت يفتح صفحة محل خشنة، السلة تتفتح...). الداتا hardcoded بسيطة جوّه الشاشات ← بعدين نرجع "
             "نلمّع ونربط الـ mock repos صح ونضيف الحالات.")
see(doc, "بدري جداً تقدر تمشي خلال التطبيق كله — كل شاشة وكل انتقال — حتى لو شكله خشن. ممتاز لاكتشاف شاشة ناقصة "
        "أو فجوة في التدفق. بيخدم هدف \"أشوف الـ UI كامل\" بشكل مباشر.")
pros(doc, ["تنقر خلال التطبيق كله من بدري جداً.", "بيكشف مشاكل التدفق والـ IA بدري."])
cons(doc, ["مفيش حاجة ملمّعة لفترة طويلة.", "إعادة شغل كتير — الـ stubs الخشنة بتتعاد كتابتها.",
           "الـ mock/Bloc/repo بيتلزق متأخر → فوضى في الكود.",
           "المعمارية تتأكد متأخر — لو النمط غلط تبقى لمست 18 شاشة."])
add_rule(doc)

# ---------- Approach C ----------
add_heading(doc, "المنهج ج) هجين (التوصية) — أثبت مرة، وسّع بأمان، لمّع", level=1)
add_par(doc, "أربع مراحل:", bold=True, space_after=3)
phases = [
    ("المرحلة 1 — الأساس", "theme + design tokens + خط Cairo + RTL، هيكل المجلدات (feature-first)، الـ DI "
     "(get_it)، الـ navigation (go_router) مع شِل الـ4 تابات، نمط الـ repository + models بأنواع (مسودة الـ schema)، "
     "ومصدر mock data واحد."),
    ("المرحلة 2 — شريحة رأسية واحدة (الرئيسية)", "نبني الرئيسية من الآخر للآخر: جريد + كروت + فلتر + Bloc + "
     "mock repo + كل الحالات + تلميع بالهوية. دي بتثبت الـ stack كامل (model ← repo ← bloc ← screen ← theme). "
     "لو فيه أي غلط في النمط نصلّحه هنا مرة واحدة."),
    ("المرحلة 3 — التوسّع بالنمط المثبَّت", "نبني باقي الشاشات كلها قابلة للوصول وشغّالة بالـ mock data، بإعادة "
     "استخدام نفس النمط والكومبوننتس المثبَّتة (مش stubs خشنة — كل شاشة بتستخدم الكروت والكومبوننتس الحقيقية من الـ "
     "design system). لإن النمط مقفول، الشاشات دي بتطلع بسرعة وبتناسق."),
    ("المرحلة 4 — مشية كاملة + تلميع نهائي", "ننقر خلال كل حاجة، نصلّح التدفق والـ IA، ونقفل."),
]
for title, body in phases:
    add_par(doc, title, bold=True, color=PRIMARY_DARK, space_after=2, space_before=4)
    add_bullet(doc, body, indent=0.3)
see(doc, "بعد المرحلة 2 شاشة واحدة مثالية بتثبت الشكل. بعد المرحلة 3 التطبيق كله قابل للتنقّل وملمّع بشكل معقول "
        "(مش خشن). وبعدين المرحلة 4 تقفل.")
add_rule(doc)

# ---------- Differences ----------
add_heading(doc, "الفروقات المهمة", level=1)
add_par(doc, "ج مقابل ب:", bold=True, color=PRIMARY_DARK, space_after=2)
add_bullet(doc, "في ب مرحلة التوسّع = stubs خشنة (إعادة شغل بعدين). في ج التوسّع بيستخدم الكومبوننتس المثبَّتة "
                "الملمّعة → إعادة شغل أقل بكتير.", indent=0.3)
add_par(doc, "ج مقابل أ:", bold=True, color=PRIMARY_DARK, space_after=2, space_before=4)
add_bullet(doc, "أ مابيديكش التطبيق كامل غير في الآخر. ج بيديك التنقّل الكامل أبكر (بعد المرحلة 3) بعد ما يثبت "
                "شاشة واحدة الأول.", indent=0.3)
add_rule(doc)

# ---------- Recommendation ----------
add_heading(doc, "✅ التوصية: المنهج ج (الهجين)", level=1)
add_par(doc,
        "إنت عايز تشوف الـ UI كامل عشان تستخرج المواصفات وبعدين تبني الباك. المنهج ج بيوصلك لـ\"UI كامل قابل "
        "للتنقّل\" أسرع من أ وأنضف من ب، وفي نفس الوقت بيحميك من إنك تكتشف غلط معماري بعد 18 شاشة لإنه بيثبت "
        "المعمارية على شاشة واحدة (الرئيسية) في البداية. باقي الشاشات بتتبني بنمط مثبَّت → بسرعة وتناسق وأقل أخطاء.",
        space_after=8)

out = r"D:\programing\wasalni\return_to_zero\مناهج_بناء_الفرونت.docx"
doc.save(out)
sys.stdout.buffer.write(b"SAVED OK\n")
